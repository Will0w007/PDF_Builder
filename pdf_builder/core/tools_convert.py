"""Conversions depuis et vers le PDF."""

from __future__ import annotations

import html
import io
import re
import shutil
import tempfile
from pathlib import Path

import pymupdf as fitz

from .external import find_ghostscript, find_soffice, office_to_pdf
from .registry import Param, ToolError, tool
from .utils import MM, PAPER_CHOICES, PAPER_SIZES, fit_rect, parse_pages, safe_filename

PAGES_HELP = "Ex. : 1-3,5,8-fin · vide = toutes"


# ---------------------------------------------------------------------------
# PDF → formats bureautiques
# ---------------------------------------------------------------------------
@tool(
    id="to_word",
    name="PDF → Word",
    category="convert",
    icon="📝",
    description="Convertit le PDF en document Word (.docx) modifiable en conservant mise en page, tableaux et images.",
    keywords="word docx doc convertir modifier",
    output_kind="files",
    params=[Param("pages", "Pages", "pages", "", help=PAGES_HELP)],
)
def to_word(path: Path, p: dict, ctx) -> list[Path]:
    try:
        from pdf2docx import Converter
    except ImportError as exc:
        raise ToolError("Le module « pdf2docx » n'est pas installé.") from exc
    doc = ctx.open(path)
    pages = parse_pages(p["pages"], doc.page_count)
    tmp = None
    src = path
    if doc.is_encrypted or path.suffix.lower() != ".pdf":
        tmp = Path(tempfile.mkdtemp()) / "src.pdf"
        doc.save(tmp, encryption=fitz.PDF_ENCRYPT_NONE)
        src = tmp
    doc.close()
    dest = ctx.output(path, "", ".docx")
    import logging

    logging.getLogger().setLevel(logging.ERROR)
    cv = Converter(str(src))
    try:
        ctx.progress(0.1, "Analyse de la mise en page…")
        cv.convert(str(dest), pages=pages)
    finally:
        cv.close()
        if tmp:
            shutil.rmtree(tmp.parent, ignore_errors=True)
    return [dest]


@tool(
    id="to_excel",
    name="PDF → Excel",
    category="convert",
    icon="📊",
    description="Détecte les tableaux du PDF et les exporte dans un classeur Excel (une feuille par tableau).",
    keywords="excel xlsx tableau tableur csv données",
    output_kind="files",
    params=[
        Param("pages", "Pages", "pages", "", help=PAGES_HELP),
        Param("layout", "Organisation", "choice", "sheet_per_table", choices=[("sheet_per_table", "Une feuille par tableau"), ("one_sheet", "Tous les tableaux sur une feuille")]),
        Param("fallback_text", "Si aucun tableau : exporter le texte ligne par ligne", "bool", True),
    ],
)
def to_excel(path: Path, p: dict, ctx) -> list[Path]:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    doc = ctx.open(path)
    wb = Workbook()
    wb.remove(wb.active)
    pages = parse_pages(p["pages"], doc.page_count)
    count = 0
    single = None
    row_cursor = 1
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="4F46E5")
    for k, pno in enumerate(pages):
        ctx.progress(k / len(pages), f"Page {pno + 1}")
        page = doc[pno]
        try:
            tables = page.find_tables().tables
        except Exception:  # noqa: BLE001
            tables = []
        for t in tables:
            rows = t.extract()
            if not rows:
                continue
            count += 1
            if p["layout"] == "one_sheet":
                if single is None:
                    single = wb.create_sheet("Tableaux")
                ws = single
                ws.cell(row=row_cursor, column=1, value=f"Page {pno + 1} – tableau {count}").font = Font(bold=True, italic=True)
                row_cursor += 1
                start = row_cursor
            else:
                ws = wb.create_sheet(f"P{pno + 1}_T{count}"[:31])
                start = 1
            for r_i, row in enumerate(rows):
                for c_i, val in enumerate(row):
                    cell = ws.cell(row=start + r_i, column=c_i + 1, value=_typed(val))
                    if r_i == 0:
                        cell.font, cell.fill = head_font, head_fill
            if p["layout"] == "one_sheet":
                row_cursor = start + len(rows) + 1
            _autofit(ws)
    if count == 0:
        if not p["fallback_text"]:
            raise ToolError("Aucun tableau détecté dans les pages sélectionnées.")
        ws = wb.create_sheet("Texte")
        r = 1
        for pno in pages:
            for line in doc[pno].get_text("text", sort=True).splitlines():
                cols = re.split(r"\s{2,}|\t", line.strip())
                for c, v in enumerate(cols):
                    ws.cell(row=r, column=c + 1, value=_typed(v))
                r += 1
        _autofit(ws)
        ctx.report("Aucun tableau structuré détecté : le texte a été exporté ligne par ligne.")
    else:
        ctx.report(f"{count} tableau(x) exporté(s).")
    dest = ctx.output(path, "", ".xlsx")
    wb.save(dest)
    doc.close()
    return [dest]


def _typed(val):
    if val is None:
        return ""
    s = str(val).strip()
    num = s.replace(" ", "").replace("\xa0", "").replace(" ", "")
    if re.fullmatch(r"-?\d+", num):
        try:
            return int(num)
        except ValueError:
            return s
    if re.fullmatch(r"-?\d+[.,]\d+", num):
        try:
            return float(num.replace(",", "."))
        except ValueError:
            return s
    return s


def _autofit(ws):
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                widths[cell.column] = max(widths.get(cell.column, 0), min(60, len(str(cell.value))))
    from openpyxl.utils import get_column_letter

    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w + 2


@tool(
    id="to_pptx",
    name="PDF → PowerPoint",
    category="convert",
    icon="📽️",
    description="Transforme chaque page en diapositive PowerPoint haute définition (avec le texte en notes).",
    keywords="powerpoint pptx présentation diapositive slides",
    output_kind="files",
    params=[
        Param("dpi", "Qualité (DPI)", "int", 150, min=72, max=400),
        Param("notes", "Copier le texte de la page dans les notes", "bool", True),
        Param("pages", "Pages", "pages", "", help=PAGES_HELP),
    ],
)
def to_pptx(path: Path, p: dict, ctx) -> list[Path]:
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError as exc:
        raise ToolError("Le module « python-pptx » n'est pas installé.") from exc
    doc = ctx.open(path)
    pages = parse_pages(p["pages"], doc.page_count)
    prs = Presentation()
    first = doc[pages[0]].rect
    emu = 12700  # 1 pt
    prs.slide_width, prs.slide_height = Emu(int(first.width * emu)), Emu(int(first.height * emu))
    blank = prs.slide_layouts[6]
    for k, pno in enumerate(pages):
        ctx.progress(k / len(pages), f"Diapositive {k + 1}")
        page = doc[pno]
        pix = page.get_pixmap(dpi=int(p["dpi"]), alpha=False)
        slide = prs.slides.add_slide(blank)
        img = io.BytesIO(pix.tobytes("png"))
        r = fit_rect(page.rect.width, page.rect.height, fitz.Rect(0, 0, first.width, first.height))
        slide.shapes.add_picture(img, Emu(int(r.x0 * emu)), Emu(int(r.y0 * emu)), Emu(int(r.width * emu)), Emu(int(r.height * emu)))
        if p["notes"]:
            text = page.get_text("text").strip()
            if text:
                slide.notes_slide.notes_text_frame.text = text
    dest = ctx.output(path, "", ".pptx")
    prs.save(dest)
    doc.close()
    return [dest]


# ---------------------------------------------------------------------------
# PDF → images / texte
# ---------------------------------------------------------------------------
@tool(
    id="to_images",
    name="PDF → Images",
    category="convert",
    icon="🖼️",
    description="Exporte les pages en images PNG, JPEG ou TIFF à la résolution souhaitée.",
    keywords="image png jpg jpeg tiff capture export photo",
    output_kind="folder",
    params=[
        Param("format", "Format", "choice", "png", choices=[("png", "PNG (sans perte)"), ("jpg", "JPEG (léger)"), ("tiff", "TIFF")]),
        Param("dpi", "Résolution (DPI)", "int", 150, min=36, max=1200),
        Param("quality", "Qualité JPEG", "int", 90, min=10, max=100, visible_if={"format": "jpg"}),
        Param("gray", "Niveaux de gris", "bool", False),
        Param("transparent", "Fond transparent", "bool", False, visible_if={"format": "png"}),
        Param("pages", "Pages", "pages", "", help=PAGES_HELP),
    ],
)
def to_images(path: Path, p: dict, ctx) -> list[Path]:
    doc = ctx.open(path)
    pages = parse_pages(p["pages"], doc.page_count)
    folder = ctx.output_dir(path, "_images")
    width = len(str(doc.page_count))
    cs = fitz.csGRAY if p["gray"] else fitz.csRGB
    outputs = []
    for k, pno in enumerate(pages):
        ctx.progress(k / len(pages), f"Page {pno + 1}")
        alpha = bool(p["transparent"]) and p["format"] == "png"
        pix = doc[pno].get_pixmap(dpi=int(p["dpi"]), colorspace=cs, alpha=alpha)
        name = folder / f"{safe_filename(path.stem)}_{pno + 1:0{width}d}.{p['format']}"
        if p["format"] == "jpg":
            pix.save(str(name), jpg_quality=int(p["quality"]))
        elif p["format"] == "tiff":
            pix.pil_save(str(name), format="TIFF", compression="tiff_deflate", dpi=(int(p["dpi"]),) * 2)
        else:
            pix.save(str(name))
        outputs.append(name)
    doc.close()
    ctx.report(f"{len(outputs)} image(s) créée(s) dans « {folder.name} ».")
    return outputs


@tool(
    id="to_text",
    name="PDF → Texte",
    category="convert",
    icon="🔤",
    description="Extrait tout le texte du PDF dans un fichier .txt (ordre de lecture naturel).",
    keywords="texte txt extraire copier contenu",
    output_kind="files",
    params=[
        Param("pages", "Pages", "pages", "", help=PAGES_HELP),
        Param("layout", "Préserver la disposition (colonnes alignées)", "bool", False),
        Param("separator", "Marquer les sauts de page", "bool", True),
    ],
)
def to_text(path: Path, p: dict, ctx) -> list[Path]:
    doc = ctx.open(path)
    chunks = []
    for pno in parse_pages(p["pages"], doc.page_count):
        page = doc[pno]
        if p["layout"]:
            text = _layout_text(page)
        else:
            text = page.get_text("text", sort=True)
        if p["separator"]:
            chunks.append(f"──────── Page {pno + 1} ────────\n")
        chunks.append(text.rstrip() + "\n\n")
    dest = ctx.output(path, "", ".txt")
    dest.write_text("".join(chunks), encoding="utf-8")
    doc.close()
    return [dest]


def _layout_text(page: fitz.Page) -> str:
    words = page.get_text("words", sort=True)
    if not words:
        return ""
    char_w = 5.0
    lines: list[list] = []
    for w in words:
        if lines and abs(lines[-1][0][3] - w[3]) < 3:
            lines[-1].append(w)
        else:
            lines.append([w])
    out = []
    for line in lines:
        line.sort(key=lambda w: w[0])
        s = ""
        for w in line:
            col = int(w[0] / char_w)
            if len(s) < col:
                s += " " * (col - len(s))
            elif s:
                s += " "
            s += w[4]
        out.append(s)
    return "\n".join(out)


@tool(
    id="to_markdown",
    name="PDF → Markdown (IA)",
    category="convert",
    icon="🤖",
    description="Convertit le PDF en Markdown structuré (titres, listes, tableaux) prêt pour ChatGPT, Claude ou un wiki.",
    keywords="markdown md ia llm chatgpt claude notion obsidian wiki",
    output_kind="files",
    params=[Param("pages", "Pages", "pages", "", help=PAGES_HELP)],
)
def to_markdown(path: Path, p: dict, ctx) -> list[Path]:
    doc = ctx.open(path)
    pages = parse_pages(p["pages"], doc.page_count)
    md = None
    try:
        import pymupdf4llm

        ctx.progress(0.1, "Analyse de la structure…")
        md = pymupdf4llm.to_markdown(doc, pages=pages, show_progress=False)
    except Exception:  # noqa: BLE001 – repli robuste
        md = None
    if not md:
        md = _simple_markdown(doc, pages)
    dest = ctx.output(path, "", ".md")
    dest.write_text(md, encoding="utf-8")
    doc.close()
    return [dest]


def _simple_markdown(doc: fitz.Document, pages: list[int]) -> str:
    sizes: dict[float, int] = {}
    for pno in pages:
        for b in doc[pno].get_text("dict")["blocks"]:
            for line in b.get("lines", []):
                for s in line["spans"]:
                    sizes[round(s["size"])] = sizes.get(round(s["size"]), 0) + len(s["text"])
    body = max(sizes, key=sizes.get) if sizes else 11
    out = []
    for pno in pages:
        for b in doc[pno].get_text("dict", sort=True)["blocks"]:
            if b.get("type") != 0:
                continue
            text = " ".join("".join(s["text"] for s in line["spans"]) for line in b["lines"]).strip()
            if not text:
                continue
            size = max(s["size"] for line in b["lines"] for s in line["spans"])
            if size >= body * 1.6:
                out.append(f"# {text}")
            elif size >= body * 1.3:
                out.append(f"## {text}")
            elif size >= body * 1.12:
                out.append(f"### {text}")
            else:
                out.append(re.sub(r"^[•●▪◦‣-]\s*", "- ", text))
        out.append("")
    return "\n\n".join(out)


@tool(
    id="to_html",
    name="PDF → HTML",
    category="convert",
    icon="🌐",
    description="Génère une page web à partir du PDF (texte et images intégrés).",
    keywords="html web page site navigateur",
    output_kind="files",
    params=[
        Param("pages", "Pages", "pages", "", help=PAGES_HELP),
        Param("mode", "Rendu", "choice", "flow", choices=[("flow", "Texte fluide (lisible, adaptatif)"), ("exact", "Mise en page fidèle")]),
    ],
)
def to_html(path: Path, p: dict, ctx) -> list[Path]:
    doc = ctx.open(path)
    parts = []
    for pno in parse_pages(p["pages"], doc.page_count):
        page = doc[pno]
        parts.append(page.get_text("xhtml" if p["mode"] == "flow" else "html"))
    title = html.escape((doc.metadata or {}).get("title") or path.stem)
    css = "body{font-family:system-ui,Segoe UI,Arial,sans-serif;max-width:900px;margin:2rem auto;padding:0 1rem;line-height:1.5;color:#222}img{max-width:100%;height:auto}"
    if p["mode"] == "exact":
        css = "body{background:#eee}div[id^=page]{background:#fff;margin:1rem auto;box-shadow:0 2px 8px #0003}"
    page_html = f"<!doctype html><html lang='fr'><head><meta charset='utf-8'><title>{title}</title><style>{css}</style></head><body>{''.join(parts)}</body></html>"
    dest = ctx.output(path, "", ".html")
    dest.write_text(page_html, encoding="utf-8")
    doc.close()
    return [dest]


# ---------------------------------------------------------------------------
# Formats → PDF
# ---------------------------------------------------------------------------
@tool(
    id="images_to_pdf",
    name="Images → PDF",
    category="convert",
    icon="🖼️",
    description="Assemble des photos ou scans (JPG, PNG, TIFF, WebP…) en un PDF, avec orientation automatique.",
    mode="all",
    accepts="images",
    keywords="image photo jpg png scan convertir assembler",
    params=[
        Param("name", "Nom du fichier", "str", "images"),
        Param("paper", "Format des pages", "choice", "auto", choices=[("auto", "Taille de l'image")] + PAPER_CHOICES),
        Param("margin", "Marge", "float", 0, min=0, max=50, suffix=" mm", visible_if={"paper": [k for k, _ in PAPER_CHOICES]}),
        Param("separate", "Un PDF par image", "bool", False),
    ],
)
def images_to_pdf(paths: list[Path], p: dict, ctx) -> list[Path]:
    from PIL import Image, ImageOps

    out = fitz.open()
    outputs = []
    for i, path in enumerate(paths):
        ctx.progress(i / len(paths), path.name)
        frames = []
        with Image.open(path) as im:
            n_frames = getattr(im, "n_frames", 1)
            for f in range(n_frames):
                im.seek(f)
                frame = ImageOps.exif_transpose(im.copy())
                dpi = im.info.get("dpi", (96, 96))
                frames.append((frame, dpi))
        for frame, dpi in frames:
            buf = io.BytesIO()
            if frame.mode in ("RGBA", "LA", "P"):
                frame.save(buf, format="PNG")
            else:
                if frame.mode not in ("RGB", "L", "CMYK"):
                    frame = frame.convert("RGB")
                frame.save(buf, format="JPEG", quality=92)
            data = buf.getvalue()
            w_px, h_px = frame.size
            xdpi = float(dpi[0] or 96) or 96
            ydpi = float(dpi[1] if len(dpi) > 1 else xdpi) or 96
            iw, ih = w_px / xdpi * 72, h_px / ydpi * 72
            if p["paper"] == "auto":
                page = out.new_page(width=iw, height=ih)
                page.insert_image(page.rect, stream=data)
            else:
                W, H = PAPER_SIZES[p["paper"]]
                if iw > ih:
                    W, H = H, W
                page = out.new_page(width=W, height=H)
                m = float(p["margin"]) * MM
                page.insert_image(fit_rect(iw, ih, fitz.Rect(m, m, W - m, H - m)), stream=data)
        if p["separate"]:
            outputs.append(ctx.save(out, ctx.output(path, "", ".pdf")))
            out.close()
            out = fitz.open()
    if not p["separate"]:
        outputs.append(ctx.save(out, ctx.output(paths[0], name=safe_filename(p["name"] or "images"))))
    out.close()
    return outputs


@tool(
    id="office_to_pdf",
    name="Office → PDF",
    category="convert",
    icon="🏢",
    description="Convertit Word, Excel, PowerPoint, OpenDocument et RTF en PDF (nécessite LibreOffice).",
    accepts="office",
    keywords="word excel powerpoint docx xlsx pptx odt libreoffice convertir",
    params=[],
)
def office_to_pdf_tool(path: Path, p: dict, ctx) -> list[Path]:
    soffice = find_soffice(ctx.settings.get("soffice"))
    if not soffice:
        raise ToolError(
            "LibreOffice est nécessaire pour convertir les documents Office.\n"
            "Installez-le gratuitement depuis https://www.libreoffice.org puis réessayez "
            "(ou indiquez son emplacement dans les Paramètres)."
        )
    tmp = Path(tempfile.mkdtemp(prefix="pdfb_office_"))
    try:
        ctx.progress(0.2, f"Conversion de {path.name} avec LibreOffice…")
        res = office_to_pdf(path, tmp, soffice)
        dest = ctx.output(path, "", ".pdf")
        shutil.move(str(res), dest)
    except Exception as exc:  # noqa: BLE001
        raise ToolError(f"Échec de la conversion de « {path.name} » : {exc}") from exc
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    return [dest]


def markdown_to_html(text: str) -> str:
    try:
        import markdown

        return markdown.markdown(text, extensions=["tables", "fenced_code"])
    except ImportError:
        pass
    out, in_list = [], False
    for line in text.splitlines():
        s = html.escape(line)
        s = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", s)
        s = re.sub(r"\*(.+?)\*", r"<i>\1</i>", s)
        s = re.sub(r"`(.+?)`", r"<code>\1</code>", s)
        m = re.match(r"^(#{1,6})\s+(.*)", s)
        li = re.match(r"^\s*[-*+]\s+(.*)", s)
        if li:
            if not in_list:
                out.append("<ul>")
                in_list = True
            out.append(f"<li>{li.group(1)}</li>")
            continue
        if in_list:
            out.append("</ul>")
            in_list = False
        if m:
            lvl = len(m.group(1))
            out.append(f"<h{lvl}>{m.group(2)}</h{lvl}>")
        elif s.strip():
            out.append(f"<p>{s}</p>")
    if in_list:
        out.append("</ul>")
    return "\n".join(out)


@tool(
    id="text_to_pdf",
    name="Texte / HTML / Markdown → PDF",
    category="convert",
    icon="📃",
    description="Met en page des fichiers .txt, .md ou .html en PDF propres et paginés.",
    accepts="text",
    keywords="texte txt markdown md html convertir",
    params=[
        Param("paper", "Format", "choice", "A4", choices=PAPER_CHOICES),
        Param("font_size", "Taille du texte", "int", 11, min=6, max=30),
        Param("margin", "Marges", "float", 20, min=5, max=60, suffix=" mm"),
    ],
)
def text_to_pdf(path: Path, p: dict, ctx) -> list[Path]:
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError:
        text = raw.decode("cp1252", errors="replace")
    ext = path.suffix.lower()
    if ext in (".html", ".htm"):
        body = text
    elif ext in (".md", ".markdown"):
        body = markdown_to_html(text)
    else:
        body = "<pre>" + html.escape(text) + "</pre>"
    css = (
        f"* {{font-family: sans-serif; font-size: {int(p['font_size'])}pt;}} "
        "pre {white-space: pre-wrap; font-family: monospace;} h1 {font-size: 1.8em;} h2 {font-size: 1.4em;} "
        "h3 {font-size: 1.2em;} code {font-family: monospace;} table {border-collapse: collapse;} td, th {border: 1px solid #999; padding: 3px;}"
    )
    W, H = PAPER_SIZES[p["paper"]]
    m = float(p["margin"]) * MM
    dest = ctx.output(path, "", ".pdf")
    story = fitz.Story(html=body, user_css=css)
    writer = fitz.DocumentWriter(str(dest))
    mediabox = fitz.Rect(0, 0, W, H)
    where = mediabox + (m, m, -m, -m)
    more = 1
    pages = 0
    while more:
        dev = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
        pages += 1
        if pages > 5000:
            break
    writer.close()
    return [dest]


@tool(
    id="rasterize",
    name="Convertir en image (anti-copie)",
    category="convert",
    icon="🧱",
    description="Transforme chaque page en image : texte non sélectionnable, rendu figé, option niveaux de gris.",
    chainable=True,
    keywords="rastériser image anti copie figer gris noir et blanc sécuriser",
    params=[
        Param("dpi", "Résolution (DPI)", "int", 150, min=50, max=600),
        Param("gray", "Niveaux de gris", "bool", False),
        Param("quality", "Qualité JPEG", "int", 85, min=20, max=100),
    ],
)
def rasterize(path: Path, p: dict, ctx) -> list[Path]:
    doc = ctx.open(path)
    out = fitz.open()
    for page in doc:
        ctx.progress(page.number / doc.page_count, f"Page {page.number + 1}")
        pix = page.get_pixmap(dpi=int(p["dpi"]), colorspace=fitz.csGRAY if p["gray"] else fitz.csRGB, alpha=False)
        np_ = out.new_page(width=page.rect.width, height=page.rect.height)
        np_.insert_image(np_.rect, stream=pix.tobytes("jpg", jpg_quality=int(p["quality"])))
    dest = ctx.save(out, ctx.output(path, "_image"))
    doc.close()
    out.close()
    return [dest]


@tool(
    id="pdfa",
    name="Convertir en PDF/A (archivage)",
    category="convert",
    icon="🏛️",
    description="Produit un PDF/A-2b ou 3b conforme à l'archivage légal à long terme (polices incorporées, couleurs normalisées).",
    chainable=True,
    keywords="pdf/a pdfa archivage légal norme iso conservation",
    params=[Param("level", "Niveau", "choice", "2", choices=[("2", "PDF/A-2b (recommandé)"), ("3", "PDF/A-3b (avec pièces jointes)"), ("1", "PDF/A-1b (ancien)")])],
)
def pdfa(path: Path, p: dict, ctx) -> list[Path]:
    import subprocess
    import sys

    gs = find_ghostscript(ctx.settings.get("ghostscript"))
    if not gs:
        raise ToolError(
            "La conversion PDF/A utilise Ghostscript (gratuit).\n"
            "• Windows : https://ghostscript.com/releases/gsdnld.html\n"
            "• macOS : brew install ghostscript\n"
            "• Linux : sudo apt install ghostscript"
        )
    doc = ctx.open(path)
    tmp = Path(tempfile.mkdtemp(prefix="pdfb_pdfa_"))
    try:
        src = tmp / "in.pdf"
        doc.save(str(src), garbage=3, deflate=True, encryption=fitz.PDF_ENCRYPT_NONE)
        doc.close()
        out = tmp / "out.pdf"
        cmd = [
            gs,
            f"-dPDFA={p['level']}",
            "-dBATCH",
            "-dNOPAUSE",
            "-dQUIET",
            "-dNOOUTERSAVE",
            "-dPDFACompatibilityPolicy=1",
            "-sColorConversionStrategy=RGB",
            "-sProcessColorModel=DeviceRGB",
            "-dEmbedAllFonts=true",
            "-sDEVICE=pdfwrite",
            f"-sOutputFile={out}",
            str(src),
        ]
        ctx.progress(0.2, "Conversion PDF/A avec Ghostscript…")
        flags = 0x08000000 if sys.platform.startswith("win") else 0
        res = subprocess.run(cmd, capture_output=True, timeout=900, creationflags=flags)
        if res.returncode != 0 or not out.exists():
            raise ToolError("Ghostscript n'a pas pu produire le PDF/A :\n" + res.stderr.decode(errors="replace")[-800:])
        dest = ctx.output(path, f"_PDFA-{p['level']}b")
        shutil.move(str(out), dest)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    ctx.report(f"{path.name} → PDF/A-{p['level']}b")
    return [dest]
