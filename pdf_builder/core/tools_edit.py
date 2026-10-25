"""Outils d'édition « en lot » : filigranes, numérotation, tampons, formulaires…"""

from __future__ import annotations

import csv
import io
import json
from pathlib import Path

import pymupdf as fitz

from .overlay import draw_text_box, image_size, image_with_opacity, stamp, text_width
from .registry import Param, ToolError, tool
from .utils import FONT_CHOICES, MM, POSITION_CHOICES, anchor_point, expand_vars, parse_color, parse_pages, safe_filename

PAGES_HELP = "Ex. : 1-3,5,8-fin · paires · impaires · vide = toutes"


def _meta(doc):
    m = doc.metadata or {}
    return m.get("title") or "", m.get("author") or ""


# ---------------------------------------------------------------------------
# Filigrane
# ---------------------------------------------------------------------------
@tool(
    id="watermark",
    name="Filigrane",
    category="edit",
    icon="💧",
    description="Ajoute un filigrane texte ou image : diagonale, mosaïque, opacité, couleur, devant ou derrière le contenu.",
    chainable=True,
    keywords="watermark marque confidentiel brouillon logo transparence",
    params=[
        Param("kind", "Type", "choice", "text", choices=[("text", "Texte"), ("image", "Image / logo")]),
        Param("text", "Texte", "str", "CONFIDENTIEL", visible_if={"kind": "text"}, help="Variables : {date}, {fichier}, {n}, {total}"),
        Param("image", "Image", "file", "", file_filter="Images (*.png *.jpg *.jpeg *.bmp *.webp)", visible_if={"kind": "image"}),
        Param("font", "Police", "choice", "hebo", choices=FONT_CHOICES, visible_if={"kind": "text"}),
        Param("size", "Taille du texte", "int", 60, min=6, max=400, visible_if={"kind": "text"}),
        Param("scale", "Largeur de l'image (% de la page)", "int", 40, min=5, max=100, visible_if={"kind": "image"}),
        Param("color", "Couleur", "color", "#D32F2F", visible_if={"kind": "text"}),
        Param("opacity", "Opacité", "int", 25, min=1, max=100, suffix=" %"),
        Param("layout", "Disposition", "choice", "diagonal", choices=[("diagonal", "Diagonale"), ("center", "Centré horizontal"), ("tile", "Mosaïque")] + [p for p in POSITION_CHOICES if p[0] != "center"]),
        Param("angle", "Angle personnalisé (°)", "int", 0, min=-180, max=180, help="0 = selon la disposition"),
        Param("behind", "Placer derrière le contenu", "bool", False),
        Param("pages", "Pages", "pages", "", help=PAGES_HELP),
    ],
)
def watermark(path: Path, p: dict, ctx) -> list[Path]:
    doc = ctx.open(path)
    color = parse_color(p["color"], (0.8, 0.1, 0.1))
    opacity = max(0.01, min(1.0, int(p["opacity"]) / 100))
    img_bytes = None
    if p["kind"] == "image":
        if not p.get("image"):
            raise ToolError("Choisissez une image pour le filigrane.")
        img_bytes = image_with_opacity(p["image"], opacity)
        iw, ih = image_size(img_bytes)
    title, author = _meta(doc)
    pages = parse_pages(p["pages"], doc.page_count)
    for k, pno in enumerate(pages):
        ctx.progress(k / len(pages), f"Page {pno + 1}")
        page = doc[pno]
        layout = p["layout"]

        def draw(op: fitz.Page):
            R = op.rect
            if img_bytes is not None:
                w = R.width * int(p["scale"]) / 100
                h = w * ih / iw
                spots = _spots(R, layout, w, h)
                for pt in spots:
                    op.insert_image(fitz.Rect(pt.x, pt.y, pt.x + w, pt.y + h), stream=img_bytes, rotate=0, keep_proportion=True)
                return
            text = expand_vars(p["text"], n=pno + 1, total=doc.page_count, filename=path.stem, title=title, author=author)
            size = float(p["size"])
            angle = float(p["angle"]) or (_diag_angle(R) if layout in ("diagonal", "tile") else 0)
            w = text_width(text, p["font"], size)
            h = size * 1.2 * len(text.split("\n"))
            if layout == "diagonal" and not p["angle"]:
                # Réduit la taille si le texte déborde de la diagonale.
                diag = (R.width**2 + R.height**2) ** 0.5 * 0.85
                if w > diag:
                    size *= diag / w
                    w, h = text_width(text, p["font"], size), size * 1.2 * len(text.split("\n"))
            for pt in _spots(R, layout, w, h):
                draw_text_box(op, text, pt, size, p["font"], color, opacity, angle)

        stamp(page, draw, overlay=not p["behind"])
    dest = ctx.save(doc, ctx.output(path, "_filigrane"))
    doc.close()
    return [dest]


def _diag_angle(R: fitz.Rect) -> float:
    import math

    return math.degrees(math.atan2(R.height, R.width))


def _spots(R: fitz.Rect, layout: str, w: float, h: float) -> list[fitz.Point]:
    if layout in ("diagonal", "center"):
        return [fitz.Point(R.x0 + (R.width - w) / 2, R.y0 + (R.height - h) / 2)]
    if layout == "tile":
        pts = []
        gap_x, gap_y = max(w * 1.3, 80), max(h * 3.5, 80)
        y = -h
        row = 0
        while y < R.height + h:
            x = -w / 2 + (row % 2) * gap_x / 2
            while x < R.width + w:
                pts.append(fitz.Point(x, y))
                x += gap_x
            y += gap_y
            row += 1
        return pts
    return [anchor_point(R, layout, w, h, 24)]


# ---------------------------------------------------------------------------
# Numérotation, en-têtes, Bates
# ---------------------------------------------------------------------------
NUMBER_FORMATS = [
    ("{n}", "1"),
    ("{n} / {total}", "1 / 10"),
    ("Page {n}", "Page 1"),
    ("Page {n} sur {total}", "Page 1 sur 10"),
    ("- {n} -", "- 1 -"),
    ("custom", "Personnalisé…"),
]


def _roman(n: int) -> str:
    vals = [(1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"), (90, "xc"), (50, "l"), (40, "xl"), (10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")]
    out = ""
    for v, s in vals:
        while n >= v:
            out += s
            n -= v
    return out


@tool(
    id="page_numbers",
    name="Numéroter les pages",
    category="edit",
    icon="🔢",
    description="Ajoute des numéros de page avec format, position, style et page de départ au choix.",
    chainable=True,
    keywords="numéros pagination numéroter pied de page",
    params=[
        Param("format", "Format", "choice", "{n} / {total}", choices=NUMBER_FORMATS),
        Param("custom", "Format personnalisé", "str", "Page {n} sur {total}", visible_if={"format": "custom"}, help="Variables : {n}, {total}, {date}, {fichier}, {titre}"),
        Param("style", "Chiffres", "choice", "arabic", choices=[("arabic", "1, 2, 3"), ("roman", "i, ii, iii"), ("ROMAN", "I, II, III")]),
        Param("position", "Position", "choice", "bottom-center", choices=POSITION_CHOICES),
        Param("start", "Premier numéro", "int", 1, min=0, max=100000),
        Param("skip", "Ne pas numéroter les N premières pages", "int", 0, min=0, max=10000),
        Param("font", "Police", "choice", "helv", choices=FONT_CHOICES),
        Param("size", "Taille", "int", 10, min=4, max=72),
        Param("color", "Couleur", "color", "#333333"),
        Param("margin", "Marge", "float", 10, min=0, max=100, suffix=" mm"),
        Param("mirror", "Alterner gauche/droite (pages paires)", "bool", False),
    ],
)
def page_numbers(path: Path, p: dict, ctx) -> list[Path]:
    doc = ctx.open(path)
    fmt = p["custom"] if p["format"] == "custom" else p["format"]
    title, author = _meta(doc)
    skip = int(p["skip"])
    numbered_total = doc.page_count - skip + int(p["start"]) - 1
    color = parse_color(p["color"])
    for page in doc:
        if page.number < skip:
            continue
        n = page.number - skip + int(p["start"])
        num = str(n)
        if p["style"] == "roman":
            num = _roman(n)
        elif p["style"] == "ROMAN":
            num = _roman(n).upper()
        text = expand_vars(fmt.replace("{n}", num).replace("{page}", num), total=numbered_total, filename=path.stem, title=title, author=author)
        pos = p["position"]
        if p["mirror"] and n % 2 == 0:
            pos = pos.replace("right", "§").replace("left", "right").replace("§", "left")

        def draw(op, text=text, pos=pos):
            size = float(p["size"])
            w = text_width(text, p["font"], size)
            pt = anchor_point(op.rect, pos, w, size * 1.2, float(p["margin"]) * MM)
            draw_text_box(op, text, pt, size, p["font"], color)

        stamp(page, draw)
    dest = ctx.save(doc, ctx.output(path, "_numéroté"))
    doc.close()
    return [dest]


@tool(
    id="header_footer",
    name="En-tête & pied de page",
    category="edit",
    icon="📰",
    description="Ajoute des textes en haut et en bas de page (gauche, centre, droite) avec variables dynamiques.",
    chainable=True,
    keywords="entête en-tête pied de page header footer date titre",
    params=[
        Param("hl", "En-tête gauche", "str", "", placeholder="ex. {titre}"),
        Param("hc", "En-tête centre", "str", ""),
        Param("hr", "En-tête droite", "str", "", placeholder="ex. {date}"),
        Param("fl", "Pied gauche", "str", "", placeholder="ex. {fichier}"),
        Param("fc", "Pied centre", "str", ""),
        Param("fr", "Pied droite", "str", "Page {n} / {total}"),
        Param("font", "Police", "choice", "helv", choices=FONT_CHOICES),
        Param("size", "Taille", "int", 9, min=4, max=48),
        Param("color", "Couleur", "color", "#555555"),
        Param("margin", "Marge", "float", 10, min=0, max=100, suffix=" mm"),
        Param("line", "Ligne de séparation", "bool", False),
        Param("pages", "Pages", "pages", "", help=PAGES_HELP),
    ],
)
def header_footer(path: Path, p: dict, ctx) -> list[Path]:
    doc = ctx.open(path)
    title, author = _meta(doc)
    color = parse_color(p["color"])
    size = float(p["size"])
    margin = float(p["margin"]) * MM
    for pno in parse_pages(p["pages"], doc.page_count):
        page = doc[pno]

        def draw(op, pno=pno):
            for key, pos in (("hl", "top-left"), ("hc", "top-center"), ("hr", "top-right"), ("fl", "bottom-left"), ("fc", "bottom-center"), ("fr", "bottom-right")):
                tmpl = p.get(key) or ""
                if not tmpl.strip():
                    continue
                text = expand_vars(tmpl, n=pno + 1, total=doc.page_count, filename=path.stem, title=title, author=author)
                w = text_width(text, p["font"], size)
                pt = anchor_point(op.rect, pos, w, size * 1.2, margin)
                draw_text_box(op, text, pt, size, p["font"], color)
            if p["line"]:
                R = op.rect
                if any((p.get(k) or "").strip() for k in ("hl", "hc", "hr")):
                    y = margin + size * 1.5
                    op.draw_line((margin, y), (R.width - margin, y), color=color, width=0.5)
                if any((p.get(k) or "").strip() for k in ("fl", "fc", "fr")):
                    y = R.height - margin - size * 1.5
                    op.draw_line((margin, y), (R.width - margin, y), color=color, width=0.5)

        stamp(page, draw)
    dest = ctx.save(doc, ctx.output(path, "_entête"))
    doc.close()
    return [dest]


@tool(
    id="bates",
    name="Numérotation Bates",
    category="edit",
    icon="⚖️",
    description="Numérotation juridique continue sur un ou plusieurs documents (ex. DOSSIER-000001).",
    mode="all",
    chainable=False,
    keywords="bates juridique avocat cote pièces numérotation continue",
    params=[
        Param("prefix", "Préfixe", "str", "PIECE-"),
        Param("suffix", "Suffixe", "str", ""),
        Param("start", "Premier numéro", "int", 1, min=0, max=10**9),
        Param("digits", "Nombre de chiffres", "int", 6, min=1, max=12),
        Param("position", "Position", "choice", "bottom-right", choices=POSITION_CHOICES),
        Param("size", "Taille", "int", 10, min=4, max=48),
        Param("color", "Couleur", "color", "#0D47A1"),
        Param("box", "Encadrer", "bool", True),
    ],
)
def bates(paths: list[Path], p: dict, ctx) -> list[Path]:
    counter = int(p["start"])
    outputs = []
    color = parse_color(p["color"])
    lines = []
    for i, path in enumerate(paths):
        ctx.progress(i / len(paths), path.name)
        doc = ctx.open(path)
        first = counter
        for page in doc:
            label = f"{p['prefix']}{counter:0{int(p['digits'])}d}{p['suffix']}"

            def draw(op, label=label):
                size = float(p["size"])
                w = text_width(label, "cobo", size)
                pt = anchor_point(op.rect, p["position"], w + 8, size * 1.2 + 6, 20)
                if p["box"]:
                    op.draw_rect(fitz.Rect(pt.x, pt.y, pt.x + w + 8, pt.y + size * 1.2 + 6), color=color, fill=(1, 1, 1), width=0.8)
                draw_text_box(op, label, fitz.Point(pt.x + 4, pt.y + 3), size, "cobo", color)

            stamp(page, draw)
            counter += 1
        dest = ctx.save(doc, ctx.output(path, "_bates"))
        lines.append(f"{path.name} : {p['prefix']}{first:0{int(p['digits'])}d} → {p['prefix']}{counter - 1:0{int(p['digits'])}d}")
        doc.close()
        outputs.append(dest)
    ctx.report("\n".join(lines))
    return outputs


# ---------------------------------------------------------------------------
# Tampons & QR codes
# ---------------------------------------------------------------------------
STAMPS = {
    "approved": ("APPROUVÉ", "#2E7D32"),
    "paid": ("PAYÉ", "#2E7D32"),
    "confidential": ("CONFIDENTIEL", "#C62828"),
    "draft": ("BROUILLON", "#6D4C41"),
    "copy": ("COPIE", "#1565C0"),
    "received": ("REÇU", "#1565C0"),
    "rejected": ("REFUSÉ", "#C62828"),
    "urgent": ("URGENT", "#E65100"),
    "custom": ("", "#6A1B9A"),
}


@tool(
    id="stamp",
    name="Tampon",
    category="edit",
    icon="🔖",
    description="Appose un tampon professionnel (Approuvé, Payé, Confidentiel…) daté, encadré et incliné.",
    chainable=True,
    keywords="tampon cachet approuvé payé reçu stamp validation",
    params=[
        Param("kind", "Tampon", "choice", "approved", choices=[(k, v[0] or "Personnalisé…") for k, v in STAMPS.items()]),
        Param("text", "Texte du tampon", "str", "VALIDÉ", visible_if={"kind": "custom"}),
        Param("color", "Couleur", "color", "#6A1B9A", visible_if={"kind": "custom"}),
        Param("subtitle", "Sous-titre", "str", "{date}", help="Ex. : {date} – Service comptable"),
        Param("position", "Position", "choice", "top-right", choices=POSITION_CHOICES),
        Param("size", "Taille", "int", 22, min=8, max=120),
        Param("angle", "Inclinaison (°)", "int", 12, min=-45, max=45),
        Param("opacity", "Opacité", "int", 85, min=10, max=100, suffix=" %"),
        Param("pages", "Pages", "pages", "1", help=PAGES_HELP),
    ],
)
def stamp_tool(path: Path, p: dict, ctx) -> list[Path]:
    doc = ctx.open(path)
    text, color_hex = STAMPS[p["kind"]]
    if p["kind"] == "custom":
        text, color_hex = p["text"] or "VALIDÉ", p["color"]
    color = parse_color(color_hex)
    title, author = _meta(doc)
    opacity = int(p["opacity"]) / 100
    for pno in parse_pages(p["pages"], doc.page_count):
        page = doc[pno]
        sub = expand_vars(p["subtitle"] or "", n=pno + 1, total=doc.page_count, filename=path.stem, title=title, author=author)

        def draw(op, sub=sub):
            size = float(p["size"])
            w = max(text_width(text, "hebo", size), text_width(sub, "helv", size * 0.45) if sub else 0) + size * 1.2
            h = size * 1.25 + (size * 0.6 if sub else 0) + size * 0.6
            tl = anchor_point(op.rect, p["position"], w, h, 36)
            box = fitz.Rect(tl.x, tl.y, tl.x + w, tl.y + h)
            pivot = fitz.Point((box.x0 + box.x1) / 2, (box.y0 + box.y1) / 2)
            morph = (pivot, fitz.Matrix(float(p["angle"])))
            shape = op.new_shape()
            shape.draw_rect(box, radius=0.12)
            shape.draw_rect(box + (3, 3, -3, -3), radius=0.1)
            shape.finish(color=color, width=2, morph=morph, stroke_opacity=opacity, fill_opacity=opacity)
            shape.commit()
            tw = text_width(text, "hebo", size)
            op.insert_text((pivot.x - tw / 2, box.y0 + size * 0.3 + size), text, fontsize=size, fontname="hebo", color=color, morph=morph, fill_opacity=opacity)
            if sub:
                ss = size * 0.45
                sw = text_width(sub, "helv", ss)
                op.insert_text((pivot.x - sw / 2, box.y1 - size * 0.45), sub, fontsize=ss, fontname="helv", color=color, morph=morph, fill_opacity=opacity)

        stamp(page, draw)
    dest = ctx.save(doc, ctx.output(path, "_tamponné"))
    doc.close()
    return [dest]


@tool(
    id="qrcode",
    name="QR code",
    category="edit",
    icon="▦",
    description="Ajoute un QR code (lien, référence, texte libre) sur les pages choisies – idéal pour le suivi documentaire.",
    chainable=True,
    keywords="qr code lien url traçabilité référence",
    params=[
        Param("data", "Contenu du QR code", "str", "https://", help="Variables : {fichier}, {n}, {total}, {date}"),
        Param("size", "Taille", "float", 25, min=8, max=150, suffix=" mm"),
        Param("position", "Position", "choice", "bottom-right", choices=POSITION_CHOICES),
        Param("caption", "Légende sous le code", "str", ""),
        Param("pages", "Pages", "pages", "1", help=PAGES_HELP),
    ],
)
def qrcode_tool(path: Path, p: dict, ctx) -> list[Path]:
    try:
        import qrcode
    except ImportError as exc:
        raise ToolError("Le module « qrcode » n'est pas installé (pip install qrcode).") from exc
    doc = ctx.open(path)
    title, author = _meta(doc)
    cache: dict[str, bytes] = {}
    for pno in parse_pages(p["pages"], doc.page_count):
        page = doc[pno]
        data = expand_vars(p["data"], n=pno + 1, total=doc.page_count, filename=path.stem, title=title, author=author)
        if data not in cache:
            img = qrcode.make(data, border=1)
            buf = io.BytesIO()
            img.save(buf, format="PNG")
            cache[data] = buf.getvalue()

        def draw(op, data=data):
            s = float(p["size"]) * MM
            cap = p.get("caption") or ""
            h = s + (10 if cap else 0)
            tl = anchor_point(op.rect, p["position"], s, h, 20)
            op.insert_image(fitz.Rect(tl.x, tl.y, tl.x + s, tl.y + s), stream=cache[data])
            if cap:
                cw = text_width(cap, "helv", 7)
                op.insert_text((tl.x + (s - cw) / 2, tl.y + s + 8), cap, fontsize=7, fontname="helv")

        stamp(page, draw)
    dest = ctx.save(doc, ctx.output(path, "_qr"))
    doc.close()
    return [dest]


# ---------------------------------------------------------------------------
# Annotations & formulaires
# ---------------------------------------------------------------------------
@tool(
    id="flatten",
    name="Aplatir",
    category="edit",
    icon="🥞",
    description="Fusionne annotations et champs de formulaire dans le contenu : plus rien n'est modifiable.",
    chainable=True,
    keywords="aplatir figer flatten formulaire annotations",
    params=[
        Param("annots", "Aplatir les annotations", "bool", True),
        Param("widgets", "Aplatir les champs de formulaire", "bool", True),
    ],
)
def flatten(path: Path, p: dict, ctx) -> list[Path]:
    doc = ctx.open(path)
    doc.bake(annots=bool(p["annots"]), widgets=bool(p["widgets"]))
    dest = ctx.save(doc, ctx.output(path, "_aplati"))
    doc.close()
    return [dest]


@tool(
    id="remove_annots",
    name="Supprimer les annotations",
    category="edit",
    icon="🧹",
    description="Retire commentaires, surlignages, dessins et/ou liens du document.",
    chainable=True,
    keywords="annotations commentaires notes surlignage supprimer liens",
    params=[
        Param("annots", "Annotations et commentaires", "bool", True),
        Param("links", "Liens", "bool", False),
        Param("widgets", "Champs de formulaire", "bool", False),
    ],
)
def remove_annots(path: Path, p: dict, ctx) -> list[Path]:
    doc = ctx.open(path)
    removed = 0
    for page in doc:
        if p["annots"]:
            for a in list(page.annots()):
                page.delete_annot(a)
                removed += 1
        if p["links"]:
            for link in page.get_links():
                page.delete_link(link)
                removed += 1
        if p["widgets"]:
            for w in list(page.widgets()):
                page.delete_widget(w)
                removed += 1
    ctx.report(f"{path.name} : {removed} élément(s) supprimé(s).")
    dest = ctx.save(doc, ctx.output(path, "_nettoyé"))
    doc.close()
    return [dest]


def read_form_fields(doc: fitz.Document) -> list[dict]:
    fields = []
    for page in doc:
        for w in page.widgets():
            fields.append(
                {
                    "name": w.field_name,
                    "type": w.field_type_string,
                    "value": w.field_value,
                    "page": page.number + 1,
                    "choices": list(w.choice_values or []),
                    "on_state": w.on_state() if w.field_type in (fitz.PDF_WIDGET_TYPE_CHECKBOX, fitz.PDF_WIDGET_TYPE_RADIOBUTTON) else None,
                }
            )
    return fields


def fill_form(doc: fitz.Document, values: dict) -> int:
    """Remplit les champs dont le nom figure dans ``values`` ; renvoie le nombre de champs remplis."""
    lowered = {str(k).strip().lower(): v for k, v in values.items()}
    filled = 0
    for page in doc:
        for w in page.widgets():
            key = (w.field_name or "").strip().lower()
            if key not in lowered:
                continue
            v = lowered[key]
            if w.field_type in (fitz.PDF_WIDGET_TYPE_CHECKBOX, fitz.PDF_WIDGET_TYPE_RADIOBUTTON):
                truthy = str(v).strip().lower() in ("1", "true", "vrai", "oui", "yes", "x", "on", str(w.on_state()).lower())
                w.field_value = w.on_state() if truthy else "Off"
            else:
                w.field_value = "" if v is None else str(v)
            w.update()
            filled += 1
    return filled


@tool(
    id="form_merge",
    name="Publipostage de formulaire",
    category="edit",
    icon="📬",
    description="Génère un PDF rempli par ligne d'un fichier CSV/Excel/JSON (colonnes = noms des champs).",
    chainable=False,
    keywords="formulaire remplir csv excel publipostage mail merge champs lot",
    output_kind="files",
    params=[
        Param("data", "Fichier de données", "file", "", file_filter="Données (*.csv *.xlsx *.json)"),
        Param("name", "Nom des fichiers", "str", "{fichier}_{ligne}", help="Utilisez {ligne} ou le nom d'une colonne entre accolades, ex. {nom}_{prenom}"),
        Param("flatten", "Aplatir les formulaires remplis", "bool", True),
    ],
)
def form_merge(path: Path, p: dict, ctx) -> list[Path]:
    rows = load_rows(p.get("data"))
    if not rows:
        raise ToolError("Le fichier de données est vide.")
    outputs = []
    for i, row in enumerate(rows, 1):
        ctx.progress(i / len(rows), f"Ligne {i}/{len(rows)}")
        doc = ctx.open(path)
        n = fill_form(doc, row)
        if n == 0 and i == 1:
            fields = ", ".join(sorted({f["name"] for f in read_form_fields(doc)})) or "aucun"
            doc.close()
            raise ToolError(f"Aucune colonne ne correspond aux champs du formulaire.\nChamps disponibles : {fields}")
        if p["flatten"]:
            doc.bake()
        name = p["name"] or "{fichier}_{ligne}"
        values = {str(k).lower(): str(v) for k, v in row.items()}
        values.update({"ligne": str(i), "fichier": path.stem})
        for k, v in values.items():
            name = name.replace("{" + k + "}", v)
        outputs.append(ctx.save(doc, ctx.output(path, name=safe_filename(name))))
        doc.close()
    return outputs


def load_rows(file: str | None) -> list[dict]:
    if not file:
        raise ToolError("Choisissez un fichier de données (CSV, Excel ou JSON).")
    f = Path(file)
    if not f.exists():
        raise ToolError(f"Fichier introuvable : {f}")
    ext = f.suffix.lower()
    if ext == ".json":
        data = json.loads(f.read_text(encoding="utf-8"))
        return data if isinstance(data, list) else [data]
    if ext == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(f, read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(it, [])]
        return [{h: ("" if v is None else v) for h, v in zip(header, r) if h} for r in it if any(v is not None for v in r)]
    raw = f.read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t") if text.strip() else csv.excel
    return list(csv.DictReader(io.StringIO(text), dialect=dialect))


@tool(
    id="form_export",
    name="Exporter les données de formulaire",
    category="extract",
    icon="📤",
    description="Exporte les valeurs des champs de formulaire en CSV ou JSON (un ou plusieurs PDF).",
    mode="all",
    keywords="formulaire export données champs csv json",
    output_kind="files",
    params=[Param("format", "Format", "choice", "csv", choices=[("csv", "CSV (Excel)"), ("json", "JSON")])],
)
def form_export(paths: list[Path], p: dict, ctx) -> list[Path]:
    records = []
    names: list[str] = []
    for path in paths:
        doc = ctx.open(path)
        rec = {"fichier": path.name}
        for f in read_form_fields(doc):
            if f["name"] not in names:
                names.append(f["name"])
            rec[f["name"]] = f["value"]
        records.append(rec)
        doc.close()
    if not names:
        raise ToolError("Aucun champ de formulaire trouvé.")
    if p["format"] == "json":
        dest = ctx.output(paths[0], "_formulaire", ".json")
        dest.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    else:
        dest = ctx.output(paths[0], "_formulaire", ".csv")
        with open(dest, "w", newline="", encoding="utf-8-sig") as fh:
            writer = csv.DictWriter(fh, fieldnames=["fichier", *names], delimiter=";")
            writer.writeheader()
            writer.writerows(records)
    ctx.report(f"{len(records)} formulaire(s), {len(names)} champ(s) exporté(s).")
    return [dest]


@tool(
    id="stamp_date",
    name="Horodatage visible",
    category="edit",
    icon="🕒",
    description="Inscrit la date et l'heure de traitement (et un texte libre) sur chaque page.",
    chainable=True,
    keywords="date heure horodatage traçabilité",
    params=[
        Param("text", "Texte", "str", "Traité le {date} à {heure}"),
        Param("position", "Position", "choice", "top-right", choices=POSITION_CHOICES),
        Param("size", "Taille", "int", 8, min=4, max=36),
        Param("color", "Couleur", "color", "#777777"),
        Param("pages", "Pages", "pages", ""),
    ],
)
def stamp_date(path: Path, p: dict, ctx) -> list[Path]:
    doc = ctx.open(path)
    color = parse_color(p["color"])
    title, author = _meta(doc)
    for pno in parse_pages(p["pages"], doc.page_count):
        text = expand_vars(p["text"], n=pno + 1, total=doc.page_count, filename=path.stem, title=title, author=author)

        def draw(op, text=text):
            size = float(p["size"])
            w = text_width(text, "helv", size)
            draw_text_box(op, text, anchor_point(op.rect, p["position"], w, size * 1.2, 14), size, "helv", color)

        stamp(doc[pno], draw)
    dest = ctx.save(doc, ctx.output(path, "_horodaté"))
    doc.close()
    return [dest]

