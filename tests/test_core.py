from __future__ import annotations

import json
import shutil

import pymupdf as fitz
import pytest

from pdf_builder.core.registry import Context, PasswordRequired, ToolError, execute, execute_workflow, load_all
from pdf_builder.core.utils import PageRangeError, parse_groups, parse_pages

from .conftest import make_form_pdf, make_image, make_scan_pdf, make_text_pdf


def run(tool, inputs, ctx, **params):
    return execute(tool, [str(i) for i in inputs], params, ctx)


def pages_of(path):
    with fitz.open(path) as d:
        return d.page_count


# ---------------------------------------------------------------- utilitaires
def test_parse_pages():
    assert parse_pages("", 5) == [0, 1, 2, 3, 4]
    assert parse_pages("1-3,5", 5) == [0, 1, 2, 4]
    assert parse_pages("3-", 5) == [2, 3, 4]
    assert parse_pages("-2", 5) == [0, 1]
    assert parse_pages("fin", 5) == [4]
    assert parse_pages("fin-1", 5) == [3]
    assert parse_pages("2-fin", 4) == [1, 2, 3]
    assert parse_pages("paires", 5) == [1, 3]
    assert parse_pages("impaires", 5) == [0, 2, 4]
    assert parse_pages("3,1,2", 5) == [2, 0, 1]
    assert parse_pages("5-3", 5) == [4, 3, 2]
    assert parse_groups("1-2; 3-fin", 5) == [[0, 1], [2, 3, 4]]
    with pytest.raises(PageRangeError):
        parse_pages("9", 5)
    with pytest.raises(PageRangeError):
        parse_pages("abc", 5)


def test_registry_is_complete():
    reg = load_all()
    assert len(reg) >= 45
    for t in reg.values():
        assert t.name and t.description and t.icon
        keys = [p.key for p in t.params]
        assert len(keys) == len(set(keys)), t.id


# ---------------------------------------------------------------- organiser
def test_merge_with_toc_and_bookmarks(sample, sample2, tmp_path, ctx):
    img = make_image(tmp_path / "photo.png")
    out = run("merge", [sample, sample2, img], ctx, toc_page=True, blank_odd=True, name="fusion")
    assert len(out) == 1
    with fitz.open(out[0]) as d:
        # 1 page de sommaire + 6 + 3 + 1 blanche (doc impair) + 1 image
        assert d.page_count == 1 + 6 + 3 + 1 + 1
        toc = d.get_toc()
        assert toc[0][1] == "Sommaire"
        assert [t[1] for t in toc if t[0] == 1][1:] == ["sample", "second", "photo"]
        assert d[0].get_links()


def test_split_modes(sample, ctx):
    assert len(run("split", [sample], ctx, mode="every", n=2)) == 3
    assert len(run("split", [sample], ctx, mode="single")) == 6
    out = run("split", [sample], ctx, mode="at", at="3,5")
    assert [pages_of(o) for o in out] == [2, 2, 2]
    out = run("split", [sample], ctx, mode="ranges", ranges="1-2; 3-fin")
    assert [pages_of(o) for o in out] == [2, 4]
    out = run("split", [sample], ctx, mode="oddeven")
    assert [pages_of(o) for o in out] == [3, 3]
    out = run("split", [sample], ctx, mode="size", size_mb=0.001)
    assert len(out) == 6


def test_split_by_bookmarks_and_blank(tmp_path, ctx):
    path = make_text_pdf(tmp_path / "b.pdf", pages=6)
    with fitz.open(path) as d:
        d.set_toc([[1, "Partie A", 1], [1, "Partie B", 4]])
        d.new_page(pno=2)  # page blanche séparatrice
        d.save(tmp_path / "b2.pdf")
    out = run("split", [tmp_path / "b2.pdf"], ctx, mode="bookmarks", level=1)
    assert len(out) == 2 and "Partie A" in out[0].name
    out = run("split", [tmp_path / "b2.pdf"], ctx, mode="blank")
    assert [pages_of(o) for o in out] == [2, 4]


def test_page_operations(sample, ctx):
    out = run("extract_pages", [sample], ctx, pages="2,1")
    assert pages_of(out[0]) == 2
    out = run("delete_pages", [sample], ctx, pages="1-2")
    assert pages_of(out[0]) == 4
    out = run("rotate", [sample], ctx, angle="90", pages="1")
    with fitz.open(out[0]) as d:
        assert d[0].rotation == 90 and d[1].rotation == 0
    out = run("reorder", [sample], ctx, mode="reverse")
    with fitz.open(out[0]) as d:
        assert "page 6" in d[0].get_text()
    out = run("insert_blank", [sample], ctx, where="after_each", count=1)
    assert pages_of(out[0]) == 12
    with pytest.raises(ToolError):
        run("delete_pages", [sample], ctx, pages="")


def test_mix_reverse(sample, sample2, ctx):
    out = run("mix", [sample, sample2], ctx, reverse_second=True)
    with fitz.open(out[0]) as d:
        assert d.page_count == 9
        assert "page 1." in d[0].get_text() and "page 3." in d[1].get_text()


def test_insert_pdf(sample, sample2, ctx):
    out = run("insert_pdf", [sample], ctx, file=str(sample2), position=2)
    with fitz.open(out[0]) as d:
        assert d.page_count == 9
        assert "page 1." in d[2].get_text()


def test_remove_blank(tmp_path, ctx):
    path = make_text_pdf(tmp_path / "x.pdf", pages=3)
    with fitz.open(path) as d:
        d.new_page(pno=1)
        d.new_page()
        d.save(tmp_path / "y.pdf")
    out = run("remove_blank", [tmp_path / "y.pdf"], ctx)
    assert pages_of(out[0]) == 3


def test_layout_tools(tmp_path, ctx):
    path = make_text_pdf(tmp_path / "r.pdf", pages=5, rotate_last=True)
    assert pages_of(run("nup", [path], ctx, per_sheet="4")[0]) == 2
    assert pages_of(run("booklet", [path], ctx)[0]) == 4  # 8 pages imposées → 2 feuilles recto/verso
    out = run("resize", [path], ctx, paper="A5")
    with fitz.open(out[0]) as d:
        assert round(d[0].rect.width) == 420
        assert d[-1].rect.width > d[-1].rect.height  # page pivotée → paysage
    out = run("crop", [path], ctx, mode="auto", padding=2)
    with fitz.open(out[0]) as d:
        assert d[0].rect.width < 595
    out = run("crop", [path], ctx, mode="manual", top=10, bottom=10, left=10, right=10)
    with fitz.open(out[0]) as d:
        assert round(d[0].rect.width) == round(595 - 20 * 72 / 25.4)
    out = run("split_half", [path], ctx, only_landscape=False)
    assert pages_of(out[0]) == 10


# ---------------------------------------------------------------- éditer
def test_watermark_text_and_image(sample, tmp_path, ctx):
    out = run("watermark", [sample], ctx, text="SECRET {n}", layout="tile")
    with fitz.open(out[0]) as d:
        assert "SECRET 1" in d[0].get_text()
    logo = make_image(tmp_path / "logo.png", size=(100, 50))
    out = run("watermark", [sample], ctx, kind="image", image=str(logo), layout="bottom-right", behind=True)
    with fitz.open(out[0]) as d:
        assert d[0].get_images(full=True) or d[0].get_xobjects()


def test_numbering_header_bates(sample, sample2, ctx):
    out = run("page_numbers", [sample], ctx, format="Page {n} sur {total}", skip=1, style="arabic")
    with fitz.open(out[0]) as d:
        assert "Page 1 sur 5" in d[1].get_text()
        assert "sur" not in d[0].get_text()
    out = run("header_footer", [sample], ctx, hl="{titre}", fr="{n}/{total}", line=True)
    with fitz.open(out[0]) as d:
        assert "Rapport" in d[0].get_text() and "1/6" in d[0].get_text()
    out = run("bates", [sample, sample2], ctx, prefix="DOS-", digits=4)
    with fitz.open(out[1]) as d:
        assert "DOS-0007" in d[0].get_text()
    assert ctx.reports


def test_stamp_qr_date(sample, ctx):
    out = run("stamp", [sample], ctx, kind="paid", pages="1")
    with fitz.open(out[0]) as d:
        assert "PAYÉ" in d[0].get_text()
    out = run("qrcode", [sample], ctx, data="https://exemple.fr/{fichier}", caption="Réf")
    assert out[0].exists()
    out = run("stamp_date", [sample], ctx)
    with fitz.open(out[0]) as d:
        assert "Traité le" in d[0].get_text()


def test_rotated_page_stamp_is_visually_placed(tmp_path, ctx):
    path = make_text_pdf(tmp_path / "rot.pdf", pages=1, rotate_last=True)
    out = run("page_numbers", [path], ctx, format="{n}", position="bottom-right")
    with fitz.open(out[0]) as d:
        page = d[0]
        rects = page.search_for("1", quads=False)
        vis = [r * page.rotation_matrix for r in rects]
        # Le numéro doit apparaître en bas à droite de la page *affichée*.
        assert any(r.x0 > page.rect.width * 0.8 and r.y0 > page.rect.height * 0.8 for r in vis)


def test_forms(tmp_path, ctx):
    form = make_form_pdf(tmp_path / "form.pdf")
    data = tmp_path / "data.csv"
    data.write_text("nom;prenom;ville;accord\nDupont;Jean;Paris;oui\nMartin;Léa;Lyon;non\n", encoding="utf-8")
    out = run("form_merge", [form], ctx, data=str(data), name="{nom}_{prenom}", flatten=False)
    assert [o.stem for o in out] == ["Dupont_Jean", "Martin_Léa"]
    with fitz.open(out[0]) as d:
        values = {w.field_name: w.field_value for w in d[0].widgets()}
        assert values["ville"] == "Paris" and values["accord"] not in ("Off", "", False)
    exported = run("form_export", out, ctx, format="json")
    records = json.loads(exported[0].read_text(encoding="utf-8"))
    assert records[1]["nom"] == "Martin"
    flat = run("flatten", [out[0]], ctx)
    with fitz.open(flat[0]) as d:
        assert not list(d[0].widgets()) and "Paris" in d[0].get_text()


def test_remove_annots(sample, ctx):
    with fitz.open(sample) as d:
        d[0].add_text_annot((100, 100), "note")
        d.save(sample.with_name("annot.pdf"))
    out = run("remove_annots", [sample.with_name("annot.pdf")], ctx)
    with fitz.open(out[0]) as d:
        assert not list(d[0].annots())


# ---------------------------------------------------------------- convertir
def test_conversions(sample, tmp_path, ctx):
    assert run("to_word", [sample], ctx, pages="1-2")[0].suffix == ".docx"
    xlsx = run("to_excel", [sample], ctx)[0]
    assert xlsx.suffix == ".xlsx"
    assert run("to_pptx", [sample], ctx, dpi=72)[0].suffix == ".pptx"
    imgs = run("to_images", [sample], ctx, format="jpg", dpi=50, pages="1-2")
    assert len(imgs) == 2 and imgs[0].suffix == ".jpg"
    assert len(run("to_images", [sample], ctx, format="tiff", dpi=40, pages="1")) == 1
    txt = run("to_text", [sample], ctx)[0]
    assert "page 1" in txt.read_text(encoding="utf-8")
    md = run("to_markdown", [sample], ctx)[0]
    assert "Chapitre" in md.read_text(encoding="utf-8")
    html = run("to_html", [sample], ctx)[0]
    assert "<html" in html.read_text(encoding="utf-8")
    ras = run("rasterize", [sample], ctx, dpi=60, gray=True)[0]
    with fitz.open(ras) as d:
        assert not d[0].get_text().strip()


def test_excel_tables(tmp_path, ctx):
    doc = fitz.open()
    page = doc.new_page()
    y0 = 100
    for r in range(4):
        for c in range(3):
            rect = fitz.Rect(72 + c * 120, y0 + r * 25, 192 + c * 120, y0 + (r + 1) * 25)
            page.draw_rect(rect, color=(0, 0, 0), width=0.8)
            page.insert_text((rect.x0 + 5, rect.y1 - 8), "Titre" if r == 0 else f"{r * 10 + c}", fontsize=10)
    path = tmp_path / "table.pdf"
    doc.save(path)
    out = run("to_excel", [path], ctx)[0]
    from openpyxl import load_workbook

    ws = load_workbook(out).worksheets[0]
    assert ws.cell(row=2, column=1).value == 10


def test_images_text_to_pdf(tmp_path, ctx):
    imgs = [make_image(tmp_path / f"i{i}.png", size=(300 + i * 100, 200)) for i in range(3)]
    out = run("images_to_pdf", imgs, ctx, paper="A4", margin=10)
    assert pages_of(out[0]) == 3
    md = tmp_path / "note.md"
    md.write_text("# Titre\n\nUn **paragraphe** avec du texte.\n\n- point 1\n- point 2\n", encoding="utf-8")
    out = run("text_to_pdf", [md], ctx)
    with fitz.open(out[0]) as d:
        assert "Titre" in d[0].get_text() and "point 2" in d[0].get_text()


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice absent")
def test_office_to_pdf(tmp_path, ctx):
    from docx import Document

    path = tmp_path / "lettre.docx"
    d = Document()
    d.add_paragraph("Bonjour depuis Word")
    d.save(path)
    out = run("office_to_pdf", [path], ctx)
    with fitz.open(out[0]) as pdf:
        assert "Bonjour" in pdf[0].get_text()


# ---------------------------------------------------------------- sécurité
def test_protect_unlock(sample, ctx):
    out = run("protect", [sample], ctx, user_pw="secret", allow_copy=False)
    with fitz.open(out[0]) as d:
        assert d.needs_pass
    with pytest.raises(PasswordRequired):
        run("compress", [out[0]], Context(out_dir=ctx.out_dir))
    unlocked = run("unlock", [out[0]], ctx, password="secret")
    with fitz.open(unlocked[0]) as d:
        assert not d.needs_pass and not d.is_encrypted
    c2 = Context(out_dir=ctx.out_dir, passwords={str(out[0]): "secret"})
    assert run("compress", [out[0]], c2)


def test_redact(sample, ctx):
    out = run("redact", [sample], ctx, terms="Contexte", p_email=True, p_phone=True, p_iban=True)
    with fitz.open(out[0]) as d:
        text = d[0].get_text()
        assert "jean.dupont@example.com" not in text
        assert "06 12 34 56 78" not in text
        assert "FR76" not in text
        assert "Contexte" not in text
        assert "Ceci est le texte" in text
    with pytest.raises(ToolError):
        run("redact", [sample], ctx)


def test_sanitize_and_metadata(sample, ctx):
    out = run("sanitize", [sample], ctx)
    with fitz.open(out[0]) as d:
        assert not (d.metadata or {}).get("author")
    out = run("metadata", [sample], ctx, title="Nouveau titre", author="Moi", layout="TwoPageLeft")
    with fitz.open(out[0]) as d:
        assert d.metadata["title"] == "Nouveau titre" and d.metadata["author"] == "Moi"


def test_certificate_sign_verify(sample, tmp_path, ctx):
    cert = run("create_certificate", [], ctx, name="Jean Test", password="1234", folder=str(tmp_path))[0]
    assert cert.suffix == ".p12"
    signed = run("sign", [sample], ctx, cert=str(cert), cert_pw="1234", reason="Test", visible=True)[0]
    ctx.reports.clear()
    run("verify_signatures", [signed], ctx)
    assert "intacte" in ctx.reports[-1] and "Jean Test" in ctx.reports[-1]
    with pytest.raises(ToolError):
        run("sign", [sample], ctx, cert=str(cert), cert_pw="mauvais")


# ---------------------------------------------------------------- optimiser
def test_compress(tmp_path, ctx):
    import io
    import random

    from PIL import Image

    doc = fitz.open()
    for _ in range(2):
        page = doc.new_page()
        img = Image.effect_noise((1600, 1600), 60).convert("RGB")
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        page.insert_image(page.rect, stream=buf.getvalue())
    path = tmp_path / "lourd.pdf"
    doc.save(path)
    random.seed(0)
    out = run("compress", [path], ctx, level="strong", linearize=True)
    assert out[0].stat().st_size < path.stat().st_size * 0.8


def test_bookmarks_repair(sample, ctx):
    out = run("bookmarks", [sample], ctx, mode="auto")
    with fitz.open(out[0]) as d:
        toc = d.get_toc()
        assert toc and toc[0][0] == 1 and "Chapitre 1" in toc[0][1]
        assert any(t[0] == 2 for t in toc)
    assert run("repair", [sample], ctx)[0].exists()


@pytest.mark.skipif(not shutil.which("tesseract"), reason="Tesseract absent")
def test_ocr(tmp_path, ctx):
    scan = make_scan_pdf(tmp_path / "scan.pdf")
    out = run("ocr", [scan], ctx, lang="fra")
    with fitz.open(out[0]) as d:
        assert "Bonjour" in d[0].get_text()


# ---------------------------------------------------------------- extraire
def test_extract_and_inspect(sample, tmp_path, ctx):
    img = make_image(tmp_path / "p.png", size=(200, 200))
    merged = run("merge", [sample, img], ctx)[0]
    assert len(run("extract_images", [merged], ctx)) == 1
    with_pj = run("attachments", [sample], ctx, mode="add", files=[str(img)])[0]
    extracted = run("attachments", [with_pj], ctx, mode="extract")
    assert extracted[0].name == "p.png"
    ctx.reports.clear()
    run("inspect", [sample], ctx)
    assert "Pages : 6" in ctx.reports[0] and "Français" in ctx.reports[0]


def test_compare(tmp_path, ctx):
    a = make_text_pdf(tmp_path / "v1.pdf", pages=2)
    with fitz.open(a) as d:
        d[0].insert_text((72, 700), "Clause ajoutée dans la version deux", fontsize=11)
        d.save(tmp_path / "v2.pdf")
    out = run("compare", [a, tmp_path / "v2.pdf"], ctx)
    with fitz.open(out[0]) as d:
        assert d.page_count == 3 and "Rapport de comparaison" in d[0].get_text()
    assert "différence" in ctx.reports[-1]


def test_search_and_jobs(sample, tmp_path):
    from pdf_builder.core.jobs import search_job

    res = search_job(Context(), folder=str(tmp_path), query="IBAN")
    assert res and res[0]["page"] == 1


def test_workflow(sample, tmp_path, ctx):
    steps = [
        {"tool": "watermark", "params": {"text": "TEST"}},
        {"tool": "page_numbers", "params": {}},
        {"tool": "compress", "params": {"level": "light"}},
    ]
    out = execute_workflow(steps, [str(sample)], ctx)
    assert out[0].name == "sample_traité.pdf"
    with fitz.open(out[0]) as d:
        assert "TEST" in d[0].get_text() and "1 / 6" in d[0].get_text()
    with pytest.raises(ToolError):
        execute_workflow([{"tool": "split", "params": {}}], [str(sample)], ctx)


def test_builtin_recipes_are_valid():
    from pdf_builder.core.recipes import BUILTIN_RECIPES

    reg = load_all()
    for r in BUILTIN_RECIPES:
        for s in r["steps"]:
            t = reg[s["tool"]]
            assert t.chainable
            assert set(s["params"]) <= {p.key for p in t.params}


def test_cli(sample, tmp_path, capsys):
    from pdf_builder.cli import main

    assert main(["liste"]) == 0
    assert main(["aide", "compress"]) == 0
    assert main(["page_numbers", str(sample), "-p", "format=Page {n}", "skip=1", "-o", str(tmp_path / "cli")]) == 0
    out = capsys.readouterr().out
    assert "✅" in out
    assert main(["recette", "e-mail", str(sample), "-o", str(tmp_path / "cli")]) == 0
    assert main(["compress", str(sample), "-p", "inconnu=1"]) == 1


@pytest.mark.skipif(not (shutil.which("gs") or shutil.which("gswin64c")), reason="Ghostscript absent")
def test_pdfa(sample, ctx):
    out = run("pdfa", [sample], ctx, level="2")[0]
    with fitz.open(out) as d:
        assert d.page_count == 6
        assert "pdfaid" in (d.get_xml_metadata() or "")
